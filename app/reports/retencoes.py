"""Relatório "Retenções e Divergências de líquido" em XLSX (spec §5.2).

Duas abas: `Retenções` (uma linha por nota, as que pedem atenção primeiro) e
`Resumo` (filtros, contagem por análise e totais).

Não é um relatório do portal, então não há paridade a manter: o layout é nosso.
O que ele existe para mostrar é o que o portal não mostra: a nota que destaca
retenções mas traz o líquido igual ao bruto.

Os valores saem de Decimal e o XLSX os grava como número de ponto flutuante (o
formato não tem célula decimal); a soma acontece antes, em Decimal.
"""

from __future__ import annotations

from collections.abc import Callable, Iterable, Mapping
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, BinaryIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.core.planilha import forcar_texto_literal
from app.domain.retencoes import (
    ROTULO_DA_ANALISE,
    LinhaRetencao,
    NotaParaRelatorio,
    TotaisRetencoes,
    montar_linhas,
    totalizar,
)

FORMATO_DINHEIRO = "#,##0.00"
FORMATO_DATA = "DD/MM/YYYY"

ROTULO_CENARIO = {
    "estrito": "Somente retenções rotuladas",
    "ampliado": "Com PIS/COFINS",
}

AVISO_PIS_COFINS = (
    "PIS e COFINS aparecem no DANFSe como 'Débito Apuração Própria' (tributo do "
    "próprio prestador), não como 'Retidos'. Por isso ficam à parte dos totais de "
    "retenção. Se foram de fato retidos é decisão de quem escritura; as colunas "
    "'com PIS/COFINS' mostram o outro cenário."
)


@dataclass(frozen=True)
class Coluna:
    titulo: str
    valor: Callable[[LinhaRetencao], Any]
    formato: str | None = None
    largura: int = 16


def _competencia(linha: LinhaRetencao) -> str | None:
    c = linha.competencia
    return f"{c.month:02d}/{c.year}" if c else None


def _sim_nao(valor: bool) -> str:
    return "Sim" if valor else "Não"


def _observacoes(linha: LinhaRetencao) -> str | None:
    a = linha.analise
    notas: list[str] = list(a.avisos)
    if a.total_retencoes_omitido:
        notas.append(
            "Total das retenções em branco ou zero na nota, apesar das retenções destacadas"
        )
    if a.descricao_diverge_do_declarado:
        notas.append("A descrição do serviço informa um líquido diferente do declarado")
    return "; ".join(notas) or None


COLUNAS: tuple[Coluna, ...] = (
    Coluna("Chave NFS-e", lambda x: x.chave_acesso, largura=52),
    Coluna("Número NFS-e", lambda x: x.numero, largura=14),
    Coluna("Data Geração", lambda x: x.data_geracao.date() if x.data_geracao else None,
           FORMATO_DATA, 13),
    Coluna("Competência", _competencia, largura=12),
    Coluna("Papel", lambda x: x.papel.capitalize(), largura=11),
    Coluna("Situação NFS-e", lambda x: x.situacao_nota, largura=13),
    Coluna("CNPJ/CPF Prestador", lambda x: x.prestador_cnpj, largura=17),
    Coluna("Nome Prestador", lambda x: x.prestador_nome, largura=34),
    Coluna("CNPJ/CPF Tomador", lambda x: x.tomador_cnpj, largura=17),
    Coluna("Nome Tomador", lambda x: x.tomador_nome, largura=34),
    Coluna("Valor do Serviço (R$)", lambda x: x.valor_servico, FORMATO_DINHEIRO, 16),
    Coluna("IRRF (R$)", lambda x: x.irrf, FORMATO_DINHEIRO, 12),
    Coluna("Contrib. Sociais Ret. (R$)", lambda x: x.contrib_sociais_retidas, FORMATO_DINHEIRO, 16),
    Coluna("Contrib. Previd. Ret. (R$)", lambda x: x.contrib_previd_retida, FORMATO_DINHEIRO, 16),
    Coluna("ISSQN Retido (R$)", lambda x: x.issqn_retido, FORMATO_DINHEIRO, 14),
    Coluna("Total Retenções Rotuladas (R$)", lambda x: x.analise.retencoes_estritas,
           FORMATO_DINHEIRO, 18),
    Coluna("PIS - Débito Apuração Própria (R$)", lambda x: x.pis_debito, FORMATO_DINHEIRO, 18),
    Coluna("COFINS - Débito Apuração Própria (R$)", lambda x: x.cofins_debito,
           FORMATO_DINHEIRO, 18),
    Coluna("Líquido Declarado na Nota (R$)", lambda x: x.analise.liquido_declarado,
           FORMATO_DINHEIRO, 18),
    Coluna("Líquido Esperado - só rotuladas (R$)", lambda x: x.analise.liquido_estrito,
           FORMATO_DINHEIRO, 20),
    Coluna("Líquido Esperado - com PIS/COFINS (R$)", lambda x: x.analise.liquido_ampliado,
           FORMATO_DINHEIRO, 20),
    Coluna("Diferença Declarado - só rotuladas (R$)", lambda x: x.analise.diferenca_estrita,
           FORMATO_DINHEIRO, 20),
    Coluna("Diferença Declarado - com PIS/COFINS (R$)", lambda x: x.analise.diferenca_ampliada,
           FORMATO_DINHEIRO, 20),
    Coluna("Líquido na Descrição (R$)", lambda x: x.analise.liquido_na_descricao,
           FORMATO_DINHEIRO, 18),
    Coluna("Cenário que a Descrição Confirma",
           lambda x: ROTULO_CENARIO.get(x.analise.cenario_da_descricao or ""), largura=24),
    Coluna("Análise", lambda x: ROTULO_DA_ANALISE[x.analise.situacao], largura=42),
    Coluna("Requer Atenção", lambda x: _sim_nao(x.analise.requer_atencao), largura=10),
    Coluna("Observações", _observacoes, largura=60),
)


def _cabecalho(aba: Worksheet) -> None:
    for celula in aba[1]:
        celula.font = Font(bold=True)
        celula.alignment = Alignment(wrap_text=True, vertical="top")


def _escrever_retencoes(aba: Worksheet, linhas: list[LinhaRetencao]) -> None:
    aba.append([c.titulo for c in COLUNAS])
    _cabecalho(aba)
    for linha in linhas:
        aba.append([c.valor(linha) for c in COLUNAS])
        # Nome e descrição vêm do XML do fornecedor: texto que parece fórmula
        # tem de sair como texto (app/core/planilha.py).
        forcar_texto_literal(aba[aba.max_row])

    for indice, coluna in enumerate(COLUNAS, start=1):
        letra = get_column_letter(indice)
        aba.column_dimensions[letra].width = coluna.largura
        if coluna.formato:
            for celula in aba[letra][1:]:
                celula.number_format = coluna.formato
    aba.freeze_panes = "A2"
    aba.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}{max(len(linhas), 1) + 1}"


def _escrever_resumo(
    aba: Worksheet,
    totais: TotaisRetencoes,
    filtros: Mapping[str, str],
    sem_competencia: int,
    gerado_em: datetime,
) -> None:
    def negrito(texto: str) -> None:
        aba.append([texto])
        aba[aba.max_row][0].font = Font(bold=True)

    negrito("Relatório de Retenções e Divergências de líquido")
    aba.append([])
    negrito("Filtros aplicados")
    for nome, valor in filtros.items():
        aba.append([nome, valor])
    if sem_competencia:
        aba.append(["Notas sem competência, fora do filtro de período", sem_competencia])
    aba.append([])

    negrito("Análise do líquido")
    aba.append(["Situação", "Qtd"])
    for situacao, rotulo in ROTULO_DA_ANALISE.items():
        aba.append([rotulo, totais.por_analise[situacao]])
    aba.append(["Total de notas", totais.notas])
    aba.append(["Notas que pedem atenção", totais.com_atencao])
    aba.append([])

    negrito("Totais das notas listadas (R$)")
    dinheiro_inicio = aba.max_row + 1
    for rotulo, valor in (
        ("Valor dos serviços", totais.valor_servico),
        ("IRRF", totais.irrf),
        ("Contribuições sociais retidas", totais.contrib_sociais_retidas),
        ("Contribuição previdenciária retida", totais.contrib_previd_retida),
        ("ISSQN retido", totais.issqn_retido),
        ("Total das retenções rotuladas", totais.retencoes_estritas),
    ):
        aba.append([rotulo, valor])
    aba.append([])
    negrito("PIS e COFINS de apuração própria (fora do total retido)")
    for rotulo, valor in (("PIS", totais.pis_debito), ("COFINS", totais.cofins_debito)):
        aba.append([rotulo, valor])
    aba.append([])
    negrito("Retenção destacada e não abatida do líquido (R$)")
    for rotulo, valor in (
        ("Só retenções rotuladas", totais.nao_abatido_estrito),
        ("Com PIS/COFINS", totais.nao_abatido_ampliado),
    ):
        aba.append([rotulo, valor])
    for linha in aba.iter_rows(min_row=dinheiro_inicio, min_col=2, max_col=2):
        for celula in linha:
            if not isinstance(celula.value, int):
                celula.number_format = FORMATO_DINHEIRO

    aba.append([])
    aba.append([AVISO_PIS_COFINS])
    aba[aba.max_row][0].alignment = Alignment(wrap_text=True, vertical="top")
    aba.append([])
    aba.append(["Gerado em:", gerado_em.strftime("%d/%m/%Y, %H:%M:%S")])

    aba.column_dimensions["A"].width = 60
    aba.column_dimensions["B"].width = 26
    for linha in aba.iter_rows():
        forcar_texto_literal(linha)


def gerar_relatorio_retencoes(
    notas: Iterable[NotaParaRelatorio],
    destino: Path | BinaryIO,
    *,
    filtros: Mapping[str, str] | None = None,
    somente_divergentes: bool = False,
    sem_competencia: int = 0,
    gerado_em: datetime | None = None,
) -> TotaisRetencoes:
    """Escreve o XLSX e devolve os totais, para o chamador e os testes."""
    linhas = montar_linhas(notas, somente_divergentes=somente_divergentes)
    totais = totalizar(linhas)

    aplicados = dict(filtros or {})
    aplicados["Somente notas com divergência"] = _sim_nao(somente_divergentes)

    livro = Workbook()
    aba_retencoes = livro.active
    assert aba_retencoes is not None
    aba_retencoes.title = "Retenções"
    _escrever_retencoes(aba_retencoes, linhas)
    _escrever_resumo(
        livro.create_sheet("Resumo"), totais, aplicados, sem_competencia,
        gerado_em or datetime.now(),
    )
    livro.save(destino)  # type: ignore[arg-type]
    return totais
