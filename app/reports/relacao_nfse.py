"""Relatório "Relação de NFS-e" — paridade com o export do portal (spec §5.1).

Duas abas, `Relação` e `Resumo`, nas mesmas colunas e na mesma ordem do
arquivo que o contador exporta hoje à mão.

Diferença deliberada em relação ao portal: os valores saem de `numeric` no
banco e `Decimal` no Python, então o total fecha em 198.227,91.
"""

from __future__ import annotations

from collections.abc import Iterable, Sequence
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, BinaryIO, Protocol

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.core.planilha import forcar_texto_literal
from app.domain.layout_relacao import COLUNAS_RELACAO, ORIGEM_NO_MODELO
from app.domain.resumo import calcular_resumo

FORMATO_DINHEIRO = "#,##0.00"
FORMATO_ALIQUOTA = "#,##0.0000"
FORMATO_DATA = "DD/MM/YYYY"

COLUNAS_DINHEIRO = frozenset(c for c in COLUNAS_RELACAO if c.endswith("(R$)"))
COLUNAS_ALIQUOTA = frozenset(c for c in COLUNAS_RELACAO if c.endswith("(%)"))


class NotaExportavel(Protocol):
    """O que o relatório lê de uma nota. Mantém o gerador fora do ORM."""

    situacao: str
    valor_servico: Decimal | None
    data_geracao: datetime | None
    competencia: date | None


def _competencia_br(valor: date | None) -> str | None:
    return f"{valor.month:02d}/{valor.year}" if valor else None


def montar_linha(nota: Any) -> list[Any]:
    """Uma nota vira 60 células, na ordem do portal."""
    linha: list[Any] = []
    for coluna in COLUNAS_RELACAO:
        if coluna == "Data Geração":
            valor = nota.data_geracao.date() if nota.data_geracao else None
        elif coluna == "Competência":
            valor = _competencia_br(nota.competencia)
        elif coluna == "Simples Nacional":
            valor = _sim_nao(nota.simples_nacional)
        elif coluna == "Retenção ISSQN":
            valor = _sim_nao(nota.issqn_retido)
        elif coluna == "Situação NFS-e":
            # O portal deixa esta coluna vazia no export de referência; a
            # situação real vai na coluna "Situação", perto do fim.
            valor = None
        elif coluna in ORIGEM_NO_MODELO:
            valor = getattr(nota, ORIGEM_NO_MODELO[coluna], None)
        else:
            valor = None
        linha.append(valor)
    return linha


def _sim_nao(valor: bool | None) -> str | None:
    if valor is None:
        return None
    return "Sim" if valor else "Não"


def _escrever_relacao(aba: Worksheet, notas: Sequence[Any]) -> None:
    aba.append(list(COLUNAS_RELACAO))
    for celula in aba[1]:
        celula.font = Font(bold=True)

    for nota in notas:
        aba.append(montar_linha(nota))
        # Nome, descrição e informações vêm do XML do fornecedor. Uma string
        # iniciada por "=" viraria fórmula no Excel do contador.
        forcar_texto_literal(aba[aba.max_row])

    for indice, coluna in enumerate(COLUNAS_RELACAO, start=1):
        letra = get_column_letter(indice)
        if coluna in COLUNAS_DINHEIRO:
            formato = FORMATO_DINHEIRO
        elif coluna in COLUNAS_ALIQUOTA:
            formato = FORMATO_ALIQUOTA
        elif coluna.startswith("Data"):
            formato = FORMATO_DATA
        else:
            continue
        for celula in aba[letra][1:]:
            celula.number_format = formato

    aba.freeze_panes = "A2"
    aba.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS_RELACAO))}{len(notas) + 1}"


def _escrever_resumo(aba: Worksheet, notas: Sequence[Any], gerado_em: datetime) -> None:
    resumo = calcular_resumo(notas)

    aba.append(["Situação", "Qtd", "Valor (R$)"])
    for celula in aba[1]:
        celula.font = Font(bold=True)

    for linha in resumo.linhas:
        aba.append([linha.situacao, linha.quantidade, linha.valor])
    aba.append([resumo.total.situacao, resumo.total.quantidade, resumo.total.valor])
    for celula in aba[aba.max_row]:
        celula.font = Font(bold=True)

    for celula in aba["C"][1:]:
        celula.number_format = FORMATO_DINHEIRO

    aba.append([])
    aba.append(["Gerado em:", gerado_em.strftime("%d/%m/%Y, %H:%M:%S")])
    aba["A1"].alignment = Alignment(horizontal="left")


def gerar_relacao(
    notas: Iterable[Any],
    destino: Path | BinaryIO,
    *,
    gerado_em: datetime | None = None,
) -> None:
    """Escreve o XLSX. `notas` sai ordenada por data de geração decrescente."""
    ordenadas = sorted(
        notas,
        key=lambda n: (n.data_geracao is not None, n.data_geracao),
        reverse=True,
    )

    livro = Workbook()
    _escrever_relacao(livro.active, ordenadas)  # type: ignore[arg-type]
    livro.active.title = "Relação"  # type: ignore[union-attr]
    _escrever_resumo(
        livro.create_sheet("Resumo"), ordenadas, gerado_em or datetime.now()
    )
    livro.save(destino)  # type: ignore[arg-type]
