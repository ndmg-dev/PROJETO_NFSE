"""Texto de terceiros em planilhas (app/core/planilha.py)."""

from __future__ import annotations

import datetime as dt
import io
import zipfile
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.core.planilha import GATILHOS_DE_FORMULA, neutralizar_formula
from app.reports.relacao_nfse import gerar_relacao
from tests.apoio_referencia import NotaDaReferencia

FORMULA = '=HYPERLINK("http://exemplo.invalido","clique aqui")'


# ------------------------------------------------------------------- CSV ---


@pytest.mark.parametrize("gatilho", GATILHOS_DE_FORMULA)
def test_prefixa_todo_gatilho(gatilho: str) -> None:
    assert neutralizar_formula(f"{gatilho}cmd") == f"'{gatilho}cmd"


@pytest.mark.parametrize(
    "texto",
    ["", "ACME LTDA", "Honorários advocatícios", "1+1 não é fórmula no meio", "a=b", " =x"],
)
def test_nao_mexe_em_texto_comum(texto: str) -> None:
    assert neutralizar_formula(texto) == texto


def test_formula_classica_de_dde() -> None:
    assert neutralizar_formula("=cmd|' /C calc'!A0").startswith("'=")


# ------------------------------------------------------------------ XLSX ---


def nota(prestador: str, descricao_ignorada: str = "") -> NotaDaReferencia:
    return NotaDaReferencia(
        chave_acesso="1" * 50,
        numero="1",
        data_geracao=dt.datetime(2026, 9, 1),
        competencia=dt.date(2026, 9, 1),
        prestador_cnpj="11222333000181",
        prestador_nome=prestador,
        tomador_cnpj="07199546000162",
        tomador_nome="TOMADOR",
        valor_servico=Decimal("100.00"),
        situacao="Normal",
    )


def gerar(*notas: NotaDaReferencia) -> bytes:
    buffer = io.BytesIO()
    gerar_relacao(notas, buffer)
    return buffer.getvalue()


def test_texto_que_parece_formula_sai_como_texto_no_xlsx() -> None:
    conteudo = gerar(nota(FORMULA))
    aba = load_workbook(io.BytesIO(conteudo))["Relação"]
    celula = next(c for c in aba[2] if c.value == FORMULA)
    assert celula.data_type == "s"


def test_xlsx_nao_contem_nenhuma_tag_de_formula() -> None:
    """Confere o XML cru: <f> é o que o Excel executa, seja qual for o tipo
    que o openpyxl infira ao reler."""
    conteudo = gerar(nota(FORMULA))
    with zipfile.ZipFile(io.BytesIO(conteudo)) as z:
        for nome in z.namelist():
            if nome.startswith("xl/worksheets/"):
                xml = z.read(nome).decode()
                assert "<f>" not in xml and "<f " not in xml, nome


def test_valor_da_celula_e_o_texto_exato_sem_apostrofo() -> None:
    """No XLSX o tipo protege; não precisa alterar o dado."""
    aba = load_workbook(io.BytesIO(gerar(nota(FORMULA))))["Relação"]
    assert FORMULA in [c.value for c in aba[2]]


def test_texto_comum_e_numeros_nao_mudam_no_xlsx() -> None:
    aba = load_workbook(io.BytesIO(gerar(nota("ACME LTDA"))))["Relação"]
    valores = [c.value for c in aba[2]]
    assert "ACME LTDA" in valores
    assert any(v == 100 or v == Decimal("100.00") for v in valores)


def test_nota_sem_texto_perigoso_nao_altera_o_tipo_das_celulas() -> None:
    aba = load_workbook(io.BytesIO(gerar(nota("ACME LTDA"))))["Relação"]
    assert all(c.data_type != "f" for c in aba[2])
