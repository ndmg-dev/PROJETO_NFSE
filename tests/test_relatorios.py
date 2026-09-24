"""Endpoints de relatório: GET /empresas/{id}/relatorio e /notas.zip.

Cobrem o que o PR do painel entregou sem teste: autenticação, isolamento entre
escritórios (spec §8), conteúdo dos arquivos e a neutralização de fórmulas em
texto vindo de fornecedor.
"""

from __future__ import annotations

import csv
import io
import json
import uuid
import zipfile
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.engine import Engine

from tests.apoio_api import CNPJ, entrar

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FORMULA = '=HYPERLINK("http://exemplo.invalido","clique aqui")'


def inserir_nota(
    engine_admin: Engine, cenario: dict[str, dict[str, Any]], lado: str, chave: str,
    **campos: Any,
) -> None:
    """Insere direto pelo dono do schema. `situacao` não tem server_default,
    então precisa vir no INSERT cru."""
    dados: dict[str, Any] = {
        "escritorio_id": cenario[lado]["escritorio"],
        "empresa_id": cenario[lado]["empresa"],
        "chave_acesso": chave,
        "papel": "tomador",
        "situacao": "Normal",
        **campos,
    }
    colunas = ", ".join(dados)
    parametros = ", ".join(f":{c}" for c in dados)
    with engine_admin.begin() as c:
        c.execute(text(f"INSERT INTO nfse ({colunas}) VALUES ({parametros})"), dados)


def ler_zip(conteudo: bytes) -> zipfile.ZipFile:
    return zipfile.ZipFile(io.BytesIO(conteudo))


CHAVE_A1, CHAVE_A2, CHAVE_B = "1" * 50, "2" * 50, "9" * 50


# --------------------------------------------------------- autenticação ---


@pytest.mark.parametrize("rota", ["relatorio", "notas.zip"])
def test_exige_token(cliente: TestClient, cenario, rota: str) -> None:  # type: ignore[no-untyped-def]
    empresa = cenario["a"]["empresa"]
    assert cliente.get(f"/empresas/{empresa}/{rota}").status_code == 401


# ------------------------------------------------------- isolamento (§8) ---


@pytest.mark.parametrize("rota", ["relatorio", "notas.zip"])
def test_nao_baixa_dados_de_outro_escritorio(  # type: ignore[no-untyped-def]
    cliente: TestClient, cenario, engine_admin, rota: str
) -> None:
    inserir_nota(engine_admin, cenario, "b", CHAVE_B)
    alvo = cenario["b"]["empresa"]
    r = cliente.get(f"/empresas/{alvo}/{rota}", headers=entrar(cliente, "a@teste.com"))
    assert r.status_code == 404, "baixou dados de outro escritório"


def test_zip_da_propria_empresa_nao_traz_nota_de_outro_escritorio(  # type: ignore[no-untyped-def]
    cliente: TestClient, cenario, engine_admin
) -> None:
    inserir_nota(engine_admin, cenario, "a", CHAVE_A1)
    inserir_nota(engine_admin, cenario, "b", CHAVE_B)
    r = cliente.get(
        f"/empresas/{cenario['a']['empresa']}/notas.zip",
        headers=entrar(cliente, "a@teste.com"),
    )
    nomes = ler_zip(r.content).namelist()
    assert f"notas/{CHAVE_A1}.json" in nomes
    assert not any(CHAVE_B in n for n in nomes)


def test_empresa_inexistente(cliente: TestClient, cenario) -> None:  # type: ignore[no-untyped-def]
    r = cliente.get(
        f"/empresas/{uuid.uuid4()}/relatorio", headers=entrar(cliente, "a@teste.com")
    )
    assert r.status_code == 404


# --------------------------------------------------------------- XLSX ---


def test_relatorio_xlsx_com_as_notas(cliente: TestClient, cenario, engine_admin) -> None:  # type: ignore[no-untyped-def]
    inserir_nota(engine_admin, cenario, "a", CHAVE_A1, valor_servico="130.00",
                 prestador_nome="LEAO FERRAMENTAS", data_geracao="2026-08-31 12:00:00+00",
                 competencia="2026-08-01")
    inserir_nota(engine_admin, cenario, "a", CHAVE_A2, valor_servico="0.00",
                 prestador_nome="CAIXA CARTOES", data_geracao="2026-08-30 12:00:00+00",
                 competencia="2026-08-01")
    r = cliente.get(
        f"/empresas/{cenario['a']['empresa']}/relatorio",
        headers=entrar(cliente, "a@teste.com"),
    )
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == XLSX
    assert f"relacao_nfse_{CNPJ}.xlsx" in r.headers["content-disposition"]

    livro = load_workbook(io.BytesIO(r.content))
    assert livro.sheetnames == ["Relação", "Resumo"]
    assert livro["Relação"].max_row == 3  # cabeçalho + 2 notas
    total = next(
        linha for linha in livro["Resumo"].iter_rows(values_only=True) if linha[0] == "Total"
    )
    assert total[1] == 2


def test_relatorio_xlsx_neutraliza_formula_do_fornecedor(  # type: ignore[no-untyped-def]
    cliente: TestClient, cenario, engine_admin
) -> None:
    inserir_nota(engine_admin, cenario, "a", CHAVE_A1, prestador_nome=FORMULA)
    r = cliente.get(
        f"/empresas/{cenario['a']['empresa']}/relatorio",
        headers=entrar(cliente, "a@teste.com"),
    )
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        xml = z.read("xl/worksheets/sheet1.xml").decode()
    assert "<f>" not in xml and "<f " not in xml, "fórmula de terceiro chegou ao XLSX"
    aba = load_workbook(io.BytesIO(r.content))["Relação"]
    assert FORMULA in [c.value for c in aba[2]]


def test_relatorio_de_empresa_sem_notas(cliente: TestClient, cenario) -> None:  # type: ignore[no-untyped-def]
    r = cliente.get(
        f"/empresas/{cenario['a']['empresa']}/relatorio",
        headers=entrar(cliente, "a@teste.com"),
    )
    assert r.status_code == 200
    assert load_workbook(io.BytesIO(r.content))["Relação"].max_row == 1


# ---------------------------------------------------------------- ZIP ---


def baixar_zip(cliente: TestClient, cenario, lado: str = "a") -> zipfile.ZipFile:  # type: ignore[no-untyped-def]
    r = cliente.get(
        f"/empresas/{cenario[lado]['empresa']}/notas.zip",
        headers=entrar(cliente, f"{lado}@teste.com"),
    )
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/zip"
    assert f"notas_{CNPJ}.zip" in r.headers["content-disposition"]
    return ler_zip(r.content)


def test_zip_tem_um_json_por_nota_e_o_csv(cliente: TestClient, cenario, engine_admin) -> None:  # type: ignore[no-untyped-def]
    inserir_nota(engine_admin, cenario, "a", CHAVE_A1, valor_servico="12000.00")
    inserir_nota(engine_admin, cenario, "a", CHAVE_A2, valor_servico="0.00")
    z = baixar_zip(cliente, cenario)
    assert sorted(z.namelist()) == sorted(
        [f"notas/{CHAVE_A1}.json", f"notas/{CHAVE_A2}.json", "notas.csv"]
    )
    dados = json.loads(z.read(f"notas/{CHAVE_A1}.json"))
    assert dados["chave_acesso"] == CHAVE_A1
    assert dados["valor_servico"] == "12000.00"  # Decimal como texto, nunca float


def test_zip_de_empresa_sem_notas(cliente: TestClient, cenario) -> None:  # type: ignore[no-untyped-def]
    assert baixar_zip(cliente, cenario).namelist() == ["LEIAME.txt"]


def test_csv_neutraliza_texto_livre_mas_nao_numeros(  # type: ignore[no-untyped-def]
    cliente: TestClient, cenario, engine_admin
) -> None:
    """Texto de fornecedor ganha apóstrofo; número negativo NÃO pode ser
    corrompido pelo prefixo."""
    inserir_nota(
        engine_admin, cenario, "a", CHAVE_A1,
        prestador_nome=FORMULA,
        descricao_servico="-cmd|' /C calc'!A0",
        informacoes_complementares="@SUM(1+1)",
        valor_servico="-5.00",
    )
    z = baixar_zip(cliente, cenario)
    linha = next(csv.DictReader(io.StringIO(z.read("notas.csv").decode())))
    assert linha["prestador_nome"] == "'" + FORMULA
    assert linha["descricao_servico"] == "'-cmd|' /C calc'!A0"
    assert linha["informacoes_complementares"] == "'@SUM(1+1)"
    assert linha["valor_servico"] == "-5.00", "prefixo corrompeu um número"


def test_json_do_zip_mantem_o_texto_original(cliente: TestClient, cenario, engine_admin) -> None:  # type: ignore[no-untyped-def]
    """JSON não é aberto por planilha: o dado sai íntegro, sem apóstrofo."""
    inserir_nota(engine_admin, cenario, "a", CHAVE_A1, prestador_nome=FORMULA)
    dados = json.loads(baixar_zip(cliente, cenario).read(f"notas/{CHAVE_A1}.json"))
    assert dados["prestador_nome"] == FORMULA
