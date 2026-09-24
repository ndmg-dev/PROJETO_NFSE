"""Projeção do XML bruto para nfse, contra Postgres real (spec §7 e §10).

O teste central percorre o caminho inteiro com o ADN mockado:
  fila do ADN -> sincronização -> dfe_bruto -> parser -> nfse -> relatório XLSX.
O mock devolve a forma DECLARADA em app/adn/contrato.py e os XMLs seguem as tags
hipotéticas do parser: provam a mecânica, não que o ADN use esses nomes.
"""

from __future__ import annotations

import io
import uuid
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.engine import Engine

from app.core.armazenamento import ArmazenamentoLocal
from app.db.base import sessao_do_escritorio
from app.db.models import DfeBruto, Nfse
from app.db.repositorio import RepositorioDfeSql, registrar_resultado_sync
from app.workers.projecao import CAMPOS_PROJETADOS, projetar_empresa
from app.workers.sincronizacao import sincronizar_empresa
from tests.apoio_api import CNPJ, entrar
from tests.apoio_xml import RETENCOES_DA_NOTA_DO_AUDIO, xml_evento, xml_nota
from tests.test_sincronizacao import ClienteFalso, doc, lote

CH1, CH2, CH3 = "1" * 50, "2" * 50, "3" * 50


@pytest.fixture
def armazem(tmp_path: Path) -> ArmazenamentoLocal:
    return ArmazenamentoLocal(tmp_path)


def sincronizar(cenario, lado, roteiro, armazem, ultimo_nsu=0):  # type: ignore[no-untyped-def]
    escritorio, empresa = cenario[lado]["escritorio"], cenario[lado]["empresa"]
    with sessao_do_escritorio(escritorio) as s:
        resultado = sincronizar_empresa(
            empresa_id=empresa, ultimo_nsu=ultimo_nsu, cliente=ClienteFalso(roteiro),
            repositorio=RepositorioDfeSql(s, escritorio), armazenamento=armazem,
        )
        registrar_resultado_sync(s, empresa, resultado.status)
    return resultado


def projetar(cenario, armazem, lado="a", **kw):  # type: ignore[no-untyped-def]
    with sessao_do_escritorio(cenario[lado]["escritorio"]) as s:
        return projetar_empresa(
            s, empresa_id=cenario[lado]["empresa"], armazenamento=armazem, **kw
        )


def consultar(engine_admin: Engine, sql: str, **params):  # type: ignore[no-untyped-def]
    """SELECT devolve as linhas; UPDATE/DELETE não devolvem e não têm o que ler."""
    with engine_admin.begin() as c:
        resultado = c.execute(text(sql), params)
        return resultado.all() if resultado.returns_rows else []


# ============================================ o caminho inteiro (a pergunta) ===


def test_do_adn_ate_o_relatorio_com_as_retencoes(  # type: ignore[no-untyped-def]
    cliente, cenario, engine_admin, armazem
) -> None:
    roteiro = [
        lote(doc(1, xml_nota(CH1, valor="12000.00", extra=RETENCOES_DA_NOTA_DO_AUDIO)),
             doc(2, xml_nota(CH2, valor="0"))),
        lote(doc(3, xml_evento(CH2), tipo="Evento")),
    ]
    r = sincronizar(cenario, "a", roteiro, armazem)
    assert r.status == "ok" and r.documentos_novos == 3

    resumo = projetar(cenario, armazem)
    assert (resumo.notas, resumo.eventos, resumo.rejeitados) == (2, 1, 0)

    # o relatório, pelo endpoint real, agora traz as retenções
    resposta = cliente.get(
        f"/empresas/{cenario['a']['empresa']}/relatorio",
        headers=entrar(cliente, "a@teste.com"),
    )
    assert resposta.status_code == 200
    livro = load_workbook(io.BytesIO(resposta.content))
    aba = livro["Relação"]
    cab = [c.value for c in aba[1]]
    linhas = {
        linha[cab.index("Chave NFS-e")]: linha
        for linha in aba.iter_rows(min_row=2, values_only=True)
    }

    nota = linhas[CH1]
    assert nota[cab.index("Valor do Serviço (R$)")] == 12000
    assert nota[cab.index("IRRF (R$)")] == 180
    assert nota[cab.index("Contrib. Sociais Ret. (R$)")] == 120
    assert nota[cab.index("PIS - Débito (R$)")] == 78
    assert nota[cab.index("COFINS - Débito (R$)")] == 360
    assert nota[cab.index("Situação")] == "Normal"

    # a nota cancelada pelo evento aparece como Cancelada no relatório
    assert linhas[CH2][cab.index("Situação")] == "Cancelada"
    resumo_aba = {
        linha[0]: linha[1]
        for linha in livro["Resumo"].iter_rows(min_row=2, max_row=5, values_only=True)
    }
    assert resumo_aba["Normal"] == 1 and resumo_aba["Cancelada"] == 1


def test_sincronizacao_grava_o_bruto_o_checkpoint_e_o_status(  # type: ignore[no-untyped-def]
    cenario, engine_admin, armazem
) -> None:
    sincronizar(cenario, "a", [lote(doc(1, xml_nota(CH1))), lote(doc(2, xml_nota(CH2)))], armazem)
    brutos = consultar(
        engine_admin,
        "SELECT nsu, status, tipo_documento FROM dfe_bruto WHERE empresa_id = :e ORDER BY nsu",
        e=cenario["a"]["empresa"],
    )
    assert [tuple(b) for b in brutos] == [(1, "pendente", "NFSe"), (2, "pendente", "NFSe")]
    empresa = consultar(engine_admin, "SELECT ultimo_nsu, ultimo_sync_status FROM empresa "
                        "WHERE id = :e", e=cenario["a"]["empresa"])[0]
    assert tuple(empresa) == (2, "ok")


def test_checkpoint_por_lote_sobrevive_a_queda_no_meio(  # type: ignore[no-untyped-def]
    cenario, engine_admin, armazem
) -> None:
    """O commit do checkpoint tem de persistir o lote 1 mesmo se o 2 falhar, e o
    lote 2 tem de conseguir gravar depois do commit (tenant reaplicado)."""
    from app.adn.cliente import AdnIndisponivel

    r = sincronizar(cenario, "a",
                    [lote(doc(1, xml_nota(CH1))), lote(doc(2, xml_nota(CH2))),
                     AdnIndisponivel("caiu")], armazem)
    assert r.status == "indisponivel"
    ultimo = consultar(engine_admin, "SELECT ultimo_nsu FROM empresa WHERE id = :e",
                       e=cenario["a"]["empresa"])[0][0]
    assert ultimo == 2


# ============================================================= papel e dados ===


def test_papel_e_decidido_pelo_cnpj_da_empresa(cenario, engine_admin, armazem) -> None:  # type: ignore[no-untyped-def]
    sincronizar(cenario, "a", [lote(
        doc(1, xml_nota(CH1, tomador=CNPJ, prestador="11222333000181")),
        doc(2, xml_nota(CH2, tomador="11222333000181", prestador=CNPJ)),
    )], armazem)
    projetar(cenario, armazem)
    papeis = {r[0]: r[1] for r in consultar(
        engine_admin, "SELECT chave_acesso, papel FROM nfse WHERE empresa_id = :e",
        e=cenario["a"]["empresa"])}
    assert papeis == {CH1: "tomador", CH2: "prestador"}


def test_valores_ficam_em_numeric_sem_ruido(cenario, engine_admin, armazem) -> None:  # type: ignore[no-untyped-def]
    sincronizar(cenario, "a", [lote(doc(1, xml_nota(
        CH1, valor="12000.00", extra=RETENCOES_DA_NOTA_DO_AUDIO)))], armazem)
    projetar(cenario, armazem)
    linha = consultar(engine_admin, "SELECT valor_servico::text, irrf::text, pis_debito::text, "
                      "valor_liquido_declarado::text FROM nfse WHERE chave_acesso = :c", c=CH1)[0]
    assert tuple(linha) == ("12000.00", "180.00", "78.00", "12000.00")


def test_todo_campo_projetado_existe_no_dto_e_no_modelo() -> None:
    from app.domain.parser import NFSeDTO

    dto_campos = set(NFSeDTO.__dataclass_fields__)
    colunas = {c.name for c in Nfse.__table__.columns}
    assert set(CAMPOS_PROJETADOS) <= dto_campos, set(CAMPOS_PROJETADOS) - dto_campos
    assert set(CAMPOS_PROJETADOS) <= colunas, set(CAMPOS_PROJETADOS) - colunas


# ============================================================ rejeição e fila ===


def test_nota_sem_valor_vai_para_erro_e_nao_para_nfse(cenario, engine_admin, armazem) -> None:  # type: ignore[no-untyped-def]
    sincronizar(cenario, "a", [lote(doc(1, xml_nota(CH1, sem_valor=True)))], armazem)
    resumo = projetar(cenario, armazem)
    assert resumo.rejeitados == 1 and resumo.notas == 0
    bruto = consultar(engine_admin, "SELECT status, erro_parse, tentativas_parse FROM dfe_bruto")[0]
    assert bruto[0] == "erro" and "valor do serviço" in bruto[1] and bruto[2] == 1
    assert consultar(engine_admin, "SELECT count(*) FROM nfse")[0][0] == 0


def test_xml_mal_formado_e_preservado_em_erro(cenario, engine_admin, armazem) -> None:  # type: ignore[no-untyped-def]
    sincronizar(cenario, "a", [lote(doc(1, b"<isto nao fecha"))], armazem)
    projetar(cenario, armazem)
    caminho, status = consultar(engine_admin, "SELECT xml_path, status FROM dfe_bruto")[0]
    assert status == "erro"
    assert Path(caminho).read_bytes() == b"<isto nao fecha", "o XML bruto tem de ser preservado"


def test_intermediario_fica_em_erro_com_o_motivo(cenario, engine_admin, armazem) -> None:  # type: ignore[no-untyped-def]
    sincronizar(cenario, "a", [lote(doc(1, xml_nota(
        CH1, tomador="11222333000181", prestador="99888777000166")))], armazem)
    projetar(cenario, armazem)
    status, motivo = consultar(engine_admin, "SELECT status, erro_parse FROM dfe_bruto")[0]
    assert status == "erro" and "intermediário" in motivo


def test_documento_ruim_nao_derruba_os_outros(cenario, engine_admin, armazem) -> None:  # type: ignore[no-untyped-def]
    sincronizar(cenario, "a", [lote(
        doc(1, xml_nota(CH1)), doc(2, b"<quebrado"), doc(3, xml_nota(CH3)))], armazem)
    resumo = projetar(cenario, armazem)
    assert (resumo.notas, resumo.rejeitados) == (2, 1)


def test_campo_maior_que_a_coluna_e_erro_do_documento_e_nao_do_lote(  # type: ignore[no-untyped-def]
    cenario, engine_admin, armazem
) -> None:
    """municipio_incidencia é String(7): valor maior é recusado pelo banco."""
    ruim = xml_nota(CH1, extra="<cMunIncid>12345678901234567890</cMunIncid>")
    sincronizar(cenario, "a", [lote(doc(1, ruim), doc(2, xml_nota(CH2)))], armazem)
    resumo = projetar(cenario, armazem)
    assert (resumo.notas, resumo.rejeitados) == (1, 1)
    motivo = consultar(engine_admin, "SELECT erro_parse FROM dfe_bruto WHERE nsu = 1")[0][0]
    assert "banco recusou" in motivo


def test_falha_de_leitura_do_xml_repete_e_depois_desiste(  # type: ignore[no-untyped-def]
    cenario, engine_admin, armazem
) -> None:
    sincronizar(cenario, "a", [lote(doc(1, xml_nota(CH1)))], armazem)
    caminho = consultar(engine_admin, "SELECT xml_path FROM dfe_bruto")[0][0]
    Path(caminho).unlink()  # some do armazenamento
    estados = []
    for _ in range(3):
        projetar(cenario, armazem, reprocessar_erros=False)
        estados.append(consultar(engine_admin, "SELECT status FROM dfe_bruto")[0][0])
    assert estados == ["pendente", "pendente", "erro"]


def test_reprocessar_erros_recupera_o_que_o_parser_novo_entende(  # type: ignore[no-untyped-def]
    cenario, engine_admin, armazem
) -> None:
    sincronizar(cenario, "a", [lote(doc(1, xml_nota(CH1, sem_valor=True)))], armazem)
    projetar(cenario, armazem)
    # ninguém reprocessa 'erro' sozinho
    assert projetar(cenario, armazem).rejeitados == 0
    consultar(engine_admin, "UPDATE dfe_bruto SET status = 'pendente'")  # simula reprocesso pedido
    assert projetar(cenario, armazem).rejeitados == 1


# ==================================================== idempotência e eventos ===


def test_projetar_duas_vezes_nao_duplica(cenario, engine_admin, armazem) -> None:  # type: ignore[no-untyped-def]
    sincronizar(cenario, "a", [lote(doc(1, xml_nota(CH1)), doc(2, xml_nota(CH2)))], armazem)
    projetar(cenario, armazem)
    consultar(engine_admin, "UPDATE dfe_bruto SET status = 'pendente'")  # força reprocesso total
    projetar(cenario, armazem)
    assert consultar(engine_admin, "SELECT count(*) FROM nfse")[0][0] == 2


def test_reprocessar_a_nota_nao_ressuscita_uma_cancelada(cenario, engine_admin, armazem) -> None:  # type: ignore[no-untyped-def]
    evento = doc(2, xml_evento(CH1), tipo="Evento")
    sincronizar(cenario, "a", [lote(doc(1, xml_nota(CH1)), evento)], armazem)
    projetar(cenario, armazem)
    assert consultar(engine_admin, "SELECT situacao FROM nfse")[0][0] == "Cancelada"
    # reprocessa só a nota: o upsert não pode devolver a situação a 'Normal'
    consultar(engine_admin, "UPDATE dfe_bruto SET status = 'pendente' WHERE nsu = 1")
    projetar(cenario, armazem)
    assert consultar(engine_admin, "SELECT situacao FROM nfse")[0][0] == "Cancelada"


def test_reprocessar_tudo_reconstroi_a_situacao_pelos_eventos(  # type: ignore[no-untyped-def]
    cenario, engine_admin, armazem
) -> None:
    """Spec §7: dá para reconstruir a projeção inteira a partir dos XMLs."""
    evento = doc(2, xml_evento(CH1), tipo="Evento")
    sincronizar(cenario, "a", [lote(doc(1, xml_nota(CH1)), evento)], armazem)
    projetar(cenario, armazem)
    consultar(engine_admin, "DELETE FROM nfse_evento")
    consultar(engine_admin, "DELETE FROM nfse")
    consultar(engine_admin, "UPDATE dfe_bruto SET status = 'pendente'")
    projetar(cenario, armazem)
    assert consultar(engine_admin, "SELECT situacao FROM nfse")[0][0] == "Cancelada"
    assert consultar(engine_admin, "SELECT count(*) FROM nfse_evento")[0][0] == 1


def test_evento_repetido_nao_duplica_o_historico(cenario, engine_admin, armazem) -> None:  # type: ignore[no-untyped-def]
    sincronizar(cenario, "a", [lote(
        doc(1, xml_nota(CH1)),
        doc(2, xml_evento(CH1), tipo="Evento"),
        doc(3, xml_evento(CH1), tipo="Evento"),  # o mesmo evento entregue de novo
    )], armazem)
    projetar(cenario, armazem)
    assert consultar(engine_admin, "SELECT count(*) FROM nfse_evento")[0][0] == 1


def test_evento_de_codigo_desconhecido_e_gravado_mas_nao_muda_a_situacao(  # type: ignore[no-untyped-def]
    cenario, engine_admin, armazem
) -> None:
    sincronizar(cenario, "a", [lote(
        doc(1, xml_nota(CH1)), doc(2, xml_evento(CH1, "999999"), tipo="Evento"))], armazem)
    projetar(cenario, armazem)
    assert consultar(engine_admin, "SELECT situacao FROM nfse")[0][0] == "Normal"
    assert consultar(engine_admin, "SELECT tipo_evento FROM nfse_evento")[0][0] == "999999"


def test_evento_orfao_fica_pendente_e_e_aplicado_quando_a_nota_chega(  # type: ignore[no-untyped-def]
    cenario, engine_admin, armazem
) -> None:
    sincronizar(cenario, "a", [lote(doc(1, xml_evento(CH1), tipo="Evento"))], armazem)
    resumo = projetar(cenario, armazem)
    assert resumo.orfaos == 1 and resumo.eventos == 0
    assert consultar(engine_admin, "SELECT status FROM dfe_bruto")[0][0] == "pendente"

    sincronizar(cenario, "a", [lote(doc(2, xml_nota(CH1)))], armazem, ultimo_nsu=1)
    projetar(cenario, armazem)
    assert consultar(engine_admin, "SELECT situacao FROM nfse")[0][0] == "Cancelada"
    status = consultar(engine_admin, "SELECT status FROM dfe_bruto WHERE nsu = 1")[0][0]
    assert status == "processado"


# ============================================================ isolamento (§8) ===


def test_projecao_de_um_escritorio_nao_toca_o_outro(cenario, engine_admin, armazem) -> None:  # type: ignore[no-untyped-def]
    sincronizar(cenario, "a", [lote(doc(1, xml_nota(CH1)))], armazem)
    sincronizar(cenario, "b", [lote(doc(1, xml_nota(CH2)))], armazem)
    projetar(cenario, armazem, lado="a")
    por_escritorio = {
        r[0]: r[1] for r in consultar(
            engine_admin, "SELECT escritorio_id, count(*) FROM nfse GROUP BY escritorio_id")
    }
    assert por_escritorio == {cenario["a"]["escritorio"]: 1}, "projetou o outro escritório"
    pendente_b = consultar(
        engine_admin,
        "SELECT status FROM dfe_bruto WHERE escritorio_id = :e",
        e=cenario["b"]["escritorio"],
    )[0][0]
    assert pendente_b == "pendente"


def test_empresa_de_outro_escritorio_nao_e_projetada(cenario, armazem) -> None:  # type: ignore[no-untyped-def]
    escritorio_a = cenario["a"]["escritorio"]
    with sessao_do_escritorio(escritorio_a) as s, pytest.raises(ValueError, match="não encontrada"):
        projetar_empresa(s, empresa_id=cenario["b"]["empresa"], armazenamento=armazem)


def test_empresa_inexistente(cenario, armazem) -> None:  # type: ignore[no-untyped-def]
    with sessao_do_escritorio(cenario["a"]["escritorio"]) as s, pytest.raises(ValueError):
        projetar_empresa(s, empresa_id=uuid.uuid4(), armazenamento=armazem)


def test_modelo_dfe_bruto_tem_o_que_o_repositorio_grava() -> None:
    colunas = {c.name for c in DfeBruto.__table__.columns}
    assert {"nsu", "tipo_documento", "xml_path", "hash_sha256", "status",
            "tentativas_parse", "erro_parse", "chave_acesso"} <= colunas
