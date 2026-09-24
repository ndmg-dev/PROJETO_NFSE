"""GET /empresas/{id}/retencoes: autenticação, isolamento, filtros e conteúdo."""

from __future__ import annotations

import io
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.apoio_api import CNPJ, entrar
from tests.test_relatorios import inserir_nota

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CH_PROBLEMA, CH_OK, CH_CANCELADA, CH_OUTRO, CH_SEM_COMP = (
    "1" * 50, "2" * 50, "3" * 50, "4" * 50, "5" * 50)

AUDIO = dict(
    valor_servico="12000.00", valor_liquido_declarado="12000.00", irrf="180.00",
    contrib_sociais_retidas="120.00", pis_debito="78.00", cofins_debito="360.00",
    issqn_retido=False, competencia="2026-05-01", data_geracao="2026-05-08 12:00:00+00",
    descricao_servico="Honorarios. Valor líquido: R$ 11.262,00",
)
OK = dict(valor_servico="500.00", valor_liquido_declarado="500.00",
          competencia="2026-06-01", data_geracao="2026-06-10 12:00:00+00")


def baixar(cliente: TestClient, cenario: Any, lado: str = "a", **params: Any):  # type: ignore[no-untyped-def]
    return cliente.get(
        f"/empresas/{cenario[lado]['empresa']}/retencoes",
        headers=entrar(cliente, f"{lado}@teste.com"),
        params=params,
    )


def chaves(resposta) -> list[str]:  # type: ignore[no-untyped-def]
    aba = load_workbook(io.BytesIO(resposta.content))["Retenções"]
    return [str(linha[0]) for linha in aba.iter_rows(min_row=2, values_only=True)]


def resumo(resposta) -> dict[str, Any]:  # type: ignore[no-untyped-def]
    aba = load_workbook(io.BytesIO(resposta.content))["Resumo"]
    return {str(r[0]): r[1] for r in aba.iter_rows(values_only=True) if r[0] is not None}


@pytest.fixture
def notas(cenario, engine_admin):  # type: ignore[no-untyped-def]
    inserir_nota(engine_admin, cenario, "a", CH_PROBLEMA, **AUDIO)
    inserir_nota(engine_admin, cenario, "a", CH_OK, **OK)
    inserir_nota(engine_admin, cenario, "a", CH_CANCELADA, situacao="Cancelada", **AUDIO)
    inserir_nota(engine_admin, cenario, "a", CH_OUTRO, papel="prestador", **OK)
    inserir_nota(engine_admin, cenario, "a", CH_SEM_COMP, valor_servico="10.00",
                 valor_liquido_declarado="10.00")


# ------------------------------------------------------- acesso e isolamento ---


def test_exige_token(cliente: TestClient, cenario) -> None:  # type: ignore[no-untyped-def]
    r = cliente.get(f"/empresas/{cenario['a']['empresa']}/retencoes")
    assert r.status_code == 401


def test_nao_baixa_de_outro_escritorio(cliente: TestClient, cenario, engine_admin) -> None:  # type: ignore[no-untyped-def]
    inserir_nota(engine_admin, cenario, "b", CH_PROBLEMA, **AUDIO)
    r = cliente.get(
        f"/empresas/{cenario['b']['empresa']}/retencoes",
        headers=entrar(cliente, "a@teste.com"),
    )
    assert r.status_code == 404


def test_nota_de_outro_escritorio_nao_entra(cliente: TestClient, cenario, engine_admin) -> None:  # type: ignore[no-untyped-def]
    inserir_nota(engine_admin, cenario, "a", CH_OK, **OK)
    inserir_nota(engine_admin, cenario, "b", CH_PROBLEMA, **AUDIO)
    assert chaves(baixar(cliente, cenario)) == [CH_OK]


# ------------------------------------------------------------------ conteúdo ---


def test_cabecalhos_http(cliente: TestClient, cenario, notas) -> None:  # type: ignore[no-untyped-def]
    r = baixar(cliente, cenario)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == XLSX
    assert f"retencoes_{CNPJ}.xlsx" in r.headers["content-disposition"]


def test_padrao_lista_so_notas_normais_e_a_com_problema_vem_primeiro(  # type: ignore[no-untyped-def]
    cliente: TestClient, cenario, notas
) -> None:
    lista = chaves(baixar(cliente, cenario))
    assert CH_CANCELADA not in lista, "retenção de nota cancelada não entra na conta"
    assert lista[0] == CH_PROBLEMA
    assert set(lista) == {CH_PROBLEMA, CH_OK, CH_OUTRO, CH_SEM_COMP}


def test_empresa_sem_notas(cliente: TestClient, cenario) -> None:  # type: ignore[no-untyped-def]
    r = baixar(cliente, cenario)
    assert r.status_code == 200 and chaves(r) == []


# ------------------------------------------------------------------- filtros ---


def test_filtro_de_situacao_inclui_as_canceladas(cliente: TestClient, cenario, notas) -> None:  # type: ignore[no-untyped-def]
    lista = chaves(baixar(cliente, cenario, situacao=["Normal", "Cancelada"]))
    assert CH_CANCELADA in lista


def test_filtro_de_papel(cliente: TestClient, cenario, notas) -> None:  # type: ignore[no-untyped-def]
    assert chaves(baixar(cliente, cenario, papel="prestador")) == [CH_OUTRO]


def test_filtro_de_periodo(cliente: TestClient, cenario, notas) -> None:  # type: ignore[no-untyped-def]
    r = baixar(cliente, cenario, competencia_de="2026-06", competencia_ate="2026-06")
    assert set(chaves(r)) == {CH_OK, CH_OUTRO}


def test_periodo_aberto_de_um_lado(cliente: TestClient, cenario, notas) -> None:  # type: ignore[no-untyped-def]
    ate_maio = set(chaves(baixar(cliente, cenario, competencia_ate="2026-05")))
    assert ate_maio == {CH_PROBLEMA}


def test_nota_sem_competencia_some_do_periodo_mas_e_contada(  # type: ignore[no-untyped-def]
    cliente: TestClient, cenario, notas
) -> None:
    r = baixar(cliente, cenario, competencia_de="2026-01", competencia_ate="2026-12")
    assert CH_SEM_COMP not in chaves(r)
    assert resumo(r)["Notas sem competência, fora do filtro de período"] == 1


def test_sem_filtro_de_periodo_a_nota_sem_competencia_entra(  # type: ignore[no-untyped-def]
    cliente: TestClient, cenario, notas
) -> None:
    r = baixar(cliente, cenario)
    assert CH_SEM_COMP in chaves(r)
    assert "Notas sem competência, fora do filtro de período" not in resumo(r)


def test_somente_divergentes(cliente: TestClient, cenario, notas) -> None:  # type: ignore[no-untyped-def]
    assert chaves(baixar(cliente, cenario, somente_divergentes="true")) == [CH_PROBLEMA]


@pytest.mark.parametrize(
    "params",
    [
        {"competencia_de": "2026-13"},
        {"competencia_de": "26-05"},
        {"competencia_ate": "maio/2026"},
        {"competencia_de": "2026-09", "competencia_ate": "2026-01"},
        {"papel": "intermediario"},
        {"situacao": "Inutilizada"},
    ],
)
def test_filtro_invalido_e_422_e_nao_relatorio_errado(  # type: ignore[no-untyped-def]
    cliente: TestClient, cenario, params: dict[str, str]
) -> None:
    assert baixar(cliente, cenario, **params).status_code == 422


def test_resumo_reflete_os_filtros_aplicados(cliente: TestClient, cenario, notas) -> None:  # type: ignore[no-untyped-def]
    r = resumo(baixar(cliente, cenario, papel="tomador", competencia_de="2026-05"))
    assert r["Papel"] == "tomador"
    assert r["Competência de"] == "2026-05"
    assert r["Competência até"] == "(sem limite)"
    assert r["Empresa (CNPJ)"] == CNPJ
