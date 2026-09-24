"""XLSX do relatório de retenções, gerado direto (sem banco)."""

from __future__ import annotations

import datetime as dt
import io
import zipfile
from typing import Any

from openpyxl import load_workbook

from app.reports.retencoes import AVISO_PIS_COFINS, COLUNAS, gerar_relatorio_retencoes
from tests.apoio_notas_ret import d, nota_consistente, nota_do_audio

FORMULA = '=HYPERLINK("http://exemplo.invalido","clique aqui")'
GERADO_EM = dt.datetime(2026, 9, 24, 10, 0, 0)


def gerar(*notas: Any, **kw: Any) -> tuple[bytes, Any]:
    buffer = io.BytesIO()
    totais = gerar_relatorio_retencoes(notas, buffer, gerado_em=GERADO_EM, **kw)
    return buffer.getvalue(), totais


def linhas_por_chave(conteudo: bytes) -> tuple[list[str], dict[str, list[Any]]]:
    aba = load_workbook(io.BytesIO(conteudo))["Retenções"]
    cab = [c.value for c in aba[1]]
    return cab, {
        str(linha[0]): list(linha) for linha in aba.iter_rows(min_row=2, values_only=True)
    }


def resumo_como_dicionario(conteudo: bytes) -> dict[str, Any]:
    aba = load_workbook(io.BytesIO(conteudo))["Resumo"]
    return {
        str(linha[0]): linha[1]
        for linha in aba.iter_rows(values_only=True)
        if linha[0] is not None and len(linha) > 1
    }


# ---------------------------------------------------------------- estrutura ---


def test_abas_e_cabecalho() -> None:
    conteudo, _ = gerar(nota_do_audio())
    livro = load_workbook(io.BytesIO(conteudo))
    assert livro.sheetnames == ["Retenções", "Resumo"]
    assert [c.value for c in livro["Retenções"][1]] == [c.titulo for c in COLUNAS]


def test_sem_notas_so_o_cabecalho() -> None:
    conteudo, totais = gerar()
    assert load_workbook(io.BytesIO(conteudo))["Retenções"].max_row == 1
    assert totais.notas == 0


# ------------------------------------------------------ a nota do áudio ---


def test_linha_da_nota_do_audio_mostra_o_problema() -> None:
    conteudo, _ = gerar(nota_do_audio())
    cab, linhas = linhas_por_chave(conteudo)
    linha = linhas["1" * 50]

    def campo(nome: str) -> Any:
        return linha[cab.index(nome)]

    assert campo("Valor do Serviço (R$)") == 12000
    assert campo("IRRF (R$)") == 180
    assert campo("Contrib. Sociais Ret. (R$)") == 120
    assert campo("Total Retenções Rotuladas (R$)") == 300
    assert campo("PIS - Débito Apuração Própria (R$)") == 78
    assert campo("COFINS - Débito Apuração Própria (R$)") == 360
    assert campo("Líquido Declarado na Nota (R$)") == 12000
    assert campo("Líquido Esperado - só rotuladas (R$)") == 11700
    assert campo("Líquido Esperado - com PIS/COFINS (R$)") == 11262
    assert campo("Diferença Declarado - só rotuladas (R$)") == 300
    assert campo("Diferença Declarado - com PIS/COFINS (R$)") == 738
    assert campo("Líquido na Descrição (R$)") == 11262
    assert campo("Cenário que a Descrição Confirma") == "Com PIS/COFINS"
    assert campo("Análise") == "Retenção destacada e não abatida do líquido"
    assert campo("Requer Atenção") == "Sim"
    assert campo("Competência") == "05/2026"
    assert campo("Papel") == "Tomador"


def test_observacoes_explicam_o_que_esta_estranho() -> None:
    cab, linhas = linhas_por_chave(gerar(nota_do_audio())[0])
    obs = linhas["1" * 50][cab.index("Observações")]
    assert "Total das retenções em branco" in obs
    assert "descrição" in obs


def test_nota_consistente_nao_pede_atencao() -> None:
    cab, linhas = linhas_por_chave(gerar(nota_consistente())[0])
    linha = linhas["2" * 50]
    assert linha[cab.index("Requer Atenção")] == "Não"
    assert linha[cab.index("Análise")] == "Consistente"
    assert linha[cab.index("Observações")] is None


def test_ordem_atencao_primeiro() -> None:
    _, linhas_ordenadas = linhas_por_chave(gerar(nota_consistente(), nota_do_audio())[0])
    assert list(linhas_ordenadas) == ["1" * 50, "2" * 50]


def test_somente_divergentes() -> None:
    _, linhas = linhas_por_chave(
        gerar(nota_consistente(), nota_do_audio(), somente_divergentes=True)[0])
    assert list(linhas) == ["1" * 50]


# -------------------------------------------------------------------- resumo ---


def test_resumo_traz_totais_e_a_contagem() -> None:
    conteudo, totais = gerar(nota_do_audio(), nota_consistente())
    r = resumo_como_dicionario(conteudo)
    assert r["Total de notas"] == 2
    assert r["Notas que pedem atenção"] == 1
    assert r["Retenção destacada e não abatida do líquido"] == 1
    assert r["Só retenções rotuladas"] == 300
    assert r["Com PIS/COFINS"] == 738
    assert r["IRRF"] == 180
    assert r["Total das retenções rotuladas"] == 300
    assert totais.nao_abatido_ampliado == d("738.00")


def test_resumo_deixa_pis_e_cofins_fora_do_total_retido() -> None:
    r = resumo_como_dicionario(gerar(nota_do_audio())[0])
    assert r["Total das retenções rotuladas"] == 300  # sem os 438 de PIS/COFINS
    assert r["PIS"] == 78 and r["COFINS"] == 360


def test_resumo_explica_a_ambiguidade_do_pis_cofins() -> None:
    r = resumo_como_dicionario(gerar(nota_do_audio())[0])
    assert AVISO_PIS_COFINS in r


def test_resumo_mostra_filtros_e_o_rodape() -> None:
    r = resumo_como_dicionario(gerar(nota_do_audio(), filtros={"Papel": "tomador"})[0])
    assert r["Papel"] == "tomador"
    assert r["Somente notas com divergência"] == "Não"
    assert r["Gerado em:"] == "24/09/2026, 10:00:00"


def test_notas_sem_competencia_so_aparecem_quando_ha() -> None:
    sem = resumo_como_dicionario(gerar(nota_do_audio(), sem_competencia=0)[0])
    com = resumo_como_dicionario(gerar(nota_do_audio(), sem_competencia=3)[0])
    assert "Notas sem competência, fora do filtro de período" not in sem
    assert com["Notas sem competência, fora do filtro de período"] == 3


# --------------------------------------------------- texto de terceiros ---


def test_texto_de_fornecedor_que_parece_formula_sai_como_texto() -> None:
    """Mesma defesa da Relação: o nome vem do XML do fornecedor."""
    conteudo, _ = gerar(nota_do_audio(prestador_nome=FORMULA, descricao_servico=FORMULA))
    with zipfile.ZipFile(io.BytesIO(conteudo)) as z:
        for nome in z.namelist():
            if nome.startswith("xl/worksheets/"):
                xml = z.read(nome).decode()
                assert "<f>" not in xml and "<f " not in xml, f"fórmula em {nome}"
    cab, linhas = linhas_por_chave(conteudo)
    assert linhas["1" * 50][cab.index("Nome Prestador")] == FORMULA
